"""
Export Utilities
Provides data export functionality to various formats
"""

import logging
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

logger = logging.getLogger(__name__)


def export_to_csv(data, headers, filename):
    """
    Export data to CSV file

    Args:
        data: List of dictionaries or tuples containing data
        headers: List of column headers
        filename: Output filename

    Returns:
        bool: True if successful, False otherwise
    """
    try:
        with open(filename, 'w', newline='', encoding='utf-8-sig') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(headers)

            for row in data:
                if isinstance(row, dict):
                    writer.writerow([row.get(h, '') for h in headers])
                else:
                    writer.writerow(row)

        logger.info(f"Data exported to CSV: {filename}")
        return True
    except Exception as e:
        logger.error(f"Failed to export to CSV: {e}")
        return False


def export_to_excel(data, headers, filename, sheet_name="Data"):
    """
    Export data to Excel file with formatting

    Args:
        data: List of dictionaries or tuples containing data
        headers: List of column headers
        filename: Output filename
        sheet_name: Name of the worksheet

    Returns:
        bool: True if successful, False otherwise
    """
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Header style
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # Write headers
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # Write data
        for row_idx, row in enumerate(data, start=2):
            if isinstance(row, dict):
                for col_idx, header in enumerate(headers, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(header, ''))
            else:
                for col_idx, value in enumerate(row, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save workbook
        wb.save(filename)
        logger.info(f"Data exported to Excel: {filename}")
        return True
    except Exception as e:
        logger.error(f"Failed to export to Excel: {e}")
        return False


def format_data_for_export(rows, column_mapping):
    """
    Format database rows for export

    Args:
        rows: List of database rows (sqlite3.Row objects)
        column_mapping: Dictionary mapping column names to display names

    Returns:
        tuple: (headers, data) formatted for export
    """
    headers = list(column_mapping.values())
    data = []

    for row in rows:
        formatted_row = {}
        for db_col, display_col in column_mapping.items():
            try:
                formatted_row[display_col] = row[db_col]
            except (KeyError, IndexError):
                formatted_row[display_col] = ''
        data.append(formatted_row)

    return headers, data
